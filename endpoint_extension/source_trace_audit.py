#!/usr/bin/env python3
from __future__ import annotations
import csv, hashlib, math
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parent
ORIG=ROOT/'04_M-Ghraam2026_ORIGINAL_UNCHANGED.xlsx'
RECON=ROOT/'Source_Reconstructed_Master_44.csv'
OUT=ROOT/'Source_to_Extension_Audit.csv'
EXPECTED_ORIG='4ff48393b15d83306e9099efa30e12ee2d797cd57a0b75c44b8f99726380becb'
EXPECTED_RECON='e1804049d7e2819e8b1228eb55db81524caaa18cb6d4745000c20d3ce609d105'

def sha256(p):
    h=hashlib.sha256()
    with open(p,'rb') as f:
        for b in iter(lambda:f.read(1<<20),b''): h.update(b)
    return h.hexdigest()

def read_rows(p):
    with open(p,encoding='utf-8-sig',newline='') as f: return list(csv.DictReader(f))

def missing(x): return x is None or (isinstance(x,float) and math.isnan(x))

def compare(src,der,tol=1e-10):
    if missing(src) and missing(der): return True,''
    if missing(src) or missing(der): return False,''
    if isinstance(src,str) or isinstance(der,str): return str(src)==str(der),''
    d=abs(float(src)-float(der)); return d<=tol,d

assert sha256(ORIG)==EXPECTED_ORIG, 'canonical workbook hash mismatch'
assert sha256(RECON)==EXPECTED_RECON, 'reconstruction hash mismatch'
wbv=load_workbook(ORIG,data_only=True,read_only=True)
wbf=load_workbook(ORIG,data_only=False,read_only=True)
wsd=wbv['dist']; wse=wbv['Enrichment Factor']
bottle=float(wbv['dist']['G53'].value)
rows=read_rows(RECON); out=[]
for r in rows:
    s=int(r['Sample']); rd=re=s+2
    def add(field,src,der,tol=1e-10,note=''):
        ok,d=compare(src,der,tol)
        out.append({'Sample':s,'Field':field,'SourceValue':src,'ReconstructionValue':der,'AbsDiff':d,'PASS':ok,'Note':note})
    add('SiteCode',wsd.cell(rd,2).value,r['SiteCode'],note='dist!B')
    add('Fine_Klf_raw',wsd.cell(rd,4).value,float(r['Fine_Klf_raw']),note='dist!D')
    add('Fine_Khf_raw',wsd.cell(rd,5).value,float(r['Fine_Khf_raw']),note='dist!E')
    if s==8:
        net=float(wse.cell(re,4).value)-bottle
        note='sample 8: Enrichment Factor gross mass minus 2.999 g bottle; dist mass cell inconsistent'
    elif s<=15:
        net=float(wsd.cell(rd,7).value); note='dist!G cached net mass (formula subtracts bottle)'
    else:
        net=float(wsd.cell(rd,7).value)-bottle; note='dist!G gross mass minus 2.999 g bottle'
    add('Fine_net_mass_g_primary',net,float(r['Fine_net_mass_g']),note=note)
    add('Fine_LF_reconstructed_from_raw',10*float(wsd.cell(rd,4).value)/net,float(r['Fine_LF_reconstructed']),tol=1e-9,note='10*Klf/primary net mass')
    add('Fine_HF_reconstructed_from_raw',10*float(wsd.cell(rd,5).value)/net,float(r['Fine_HF_reconstructed']),tol=1e-9,note='10*Khf/primary net mass')
    for col,field in [(11,'Coarse_Klf_raw'),(12,'Coarse_Khf_raw'),(13,'Coarse_mass_g')]:
        src=wsd.cell(rd,col).value; der=None if r[field]=='' else float(r[field]); add(field,src,der,note=f'dist column {col}')
    if wsd.cell(rd,11).value is not None and wsd.cell(rd,13).value not in (None,0):
        add('Coarse_LF_reconstructed_from_raw',10*float(wsd.cell(rd,11).value)/float(wsd.cell(rd,13).value),float(r['Coarse_LF_reconstructed']),tol=1e-9,note='10*Klf/coarse mass')
        add('Coarse_HF_reconstructed_from_raw',10*float(wsd.cell(rd,12).value)/float(wsd.cell(rd,13).value),float(r['Coarse_HF_reconstructed']),tol=1e-9,note='10*Khf/coarse mass')
    for j,m in enumerate(['Fe','Zn','Cu','Cr','Co','Ni','Mn','Pb','Cd'],start=20):
        add(f'{m}_mgkg',wse.cell(re,j).value,float(r[f'{m}_mgkg']),tol=1e-8,note=f'Enrichment Factor final mg/kg column {j}')
with open(OUT,'w',encoding='utf-8-sig',newline='') as f:
    w=csv.DictWriter(f,fieldnames=list(out[0])); w.writeheader(); w.writerows(out)
fails=[x for x in out if not x['PASS']]
nums=[float(x['AbsDiff']) for x in out if x['AbsDiff']!='']
if fails: raise SystemExit(f'FAIL: {len(fails)} source-trace mismatches')
print(f'PASS {len(out)}/{len(out)} source-to-extension checks; max abs numeric difference {max(nums):.3e}')
