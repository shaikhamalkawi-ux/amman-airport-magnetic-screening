#!/usr/bin/env python3
"""New release utility: reconstruct the 44-row master from the unchanged workbook.

This is an independently authored release utility, not the historical author
SourceLock V3 generator. It preserves raw coordinate strings and implements the
cell map and mass rules verified against the original author source trace.
The optional --reference comparison verifies all fields without using the
reference values to construct the output. Never replace the hash-locked author
CSV merely because another CSV is numerically equivalent.
"""
from __future__ import annotations
import argparse
import csv
import hashlib
import json
import math
from pathlib import Path
import re
from openpyxl import load_workbook

WORKBOOK_SHA256 = '4ff48393b15d83306e9099efa30e12ee2d797cd57a0b75c44b8f99726380becb'
METALS = ('Fe', 'Zn', 'Cu', 'Cr', 'Co', 'Ni', 'Mn', 'Pb', 'Cd')
CATEGORICAL_FIELDS = {'SiteCode', 'Side', 'Distance_raw', 'Latitude_raw',
                      'Longitude_raw', 'Fine_mass_rule'}


def sha256(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def load_csv(path):
    with path.open(encoding='utf-8-sig', newline='') as stream:
        reader = csv.DictReader(stream)
        return reader.fieldnames, list(reader)


def reconstruct(workbook):
    book = load_workbook(workbook, data_only=True, read_only=True)
    # Read cached source values once; do not recalculate or rewrite the workbook.
    sheets = {name: list(book[name].iter_rows(values_only=True))
              for name in ('dist', 'Enrichment Factor', 'LOCATION')}
    book.close()
    def cell(sheet, row, column):
        return sheets[sheet][row - 1][column - 1]
    bottle = float(cell('dist', 53, 7))
    if bottle != 2.999:
        raise ValueError('Unexpected source container mass')
    locations = {str(cell('LOCATION', row, 2)):
                 (cell('LOCATION', row, 1), str(cell('LOCATION', row, 3)),
                  str(cell('LOCATION', row, 4))) for row in range(2, 21)}
    result = []
    for sample in range(1, 45):
        row = sample + 2
        site = str(cell('dist', row, 2))
        station = int(re.search(r'(\d+)$', site).group(1))
        distance_raw = str(cell('dist', row, 3))
        number = re.search(r'\d+(?:\.\d+)?', distance_raw)
        if number is None:
            raise ValueError(f'No numeric distance in sample {sample}')
        location_index, latitude_raw, longitude_raw = locations[site]
        klf, khf = float(cell('dist', row, 4)), float(cell('dist', row, 5))
        if sample == 8:
            net = float(cell('Enrichment Factor', row, 4)) - bottle
            rule = 'EF gross mass - 2.999 g bottle; dist mass cell inconsistent'
        elif sample <= 15:
            net = float(cell('dist', row, 7))
            rule = 'dist mass used directly (net)'
        else:
            net = float(cell('dist', row, 7)) - bottle
            rule = 'dist gross mass - 2.999 g bottle'
        lf, hf = 10 * klf / net, 10 * khf / net
        cklf, ckhf, mass = [cell('dist', row, col) for col in (11, 12, 13)]
        if cklf is None or mass in (None, 0):
            clf = chf = cfd = ''
        else:
            clf = 10 * float(cklf) / float(mass)
            chf = 10 * float(ckhf) / float(mass)
            cfd = 100 * (clf - chf) / clf
        record = {
            'Sample': sample, 'SiteCode': site, 'Side': site[0],
            'StationNumber': station, 'Distance_raw': distance_raw,
            'Distance_numeric_m': float(number.group()),
            'LocationIndex': location_index, 'Latitude_raw': latitude_raw,
            'Longitude_raw': longitude_raw,
            'Fine_Klf_raw': klf, 'Fine_Khf_raw': khf,
            'Fine_net_mass_g': net, 'Fine_mass_rule': rule,
            'Fine_LF_reconstructed': lf, 'Fine_HF_reconstructed': hf,
            'Fine_FD_pct': 100 * (lf - hf) / lf,
            'Coarse_Klf_raw': cklf, 'Coarse_Khf_raw': ckhf,
            'Coarse_mass_g': mass, 'Coarse_LF_reconstructed': clf,
            'Coarse_HF_reconstructed': chf, 'Coarse_FD_pct': cfd,
        }
        for column, metal in enumerate(METALS, 20):
            record[metal + '_mgkg'] = cell('Enrichment Factor', row, column)
        result.append(record)
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('workbook', type=Path)
    parser.add_argument('output_csv', type=Path)
    parser.add_argument('--reference', type=Path)
    parser.add_argument('--report', type=Path, default=Path('reconstruction_check.json'))
    args = parser.parse_args()
    workbook = args.workbook.resolve()
    output = args.output_csv.resolve()
    report = args.report.resolve()
    reference = args.reference.resolve() if args.reference else None
    if output.exists():
        parser.error('Output already exists; choose a new file to preserve originals.')
    if report in {workbook, output, reference}:
        parser.error('Report must not replace an input or output CSV.')
    if sha256(workbook) != WORKBOOK_SHA256:
        parser.error('Workbook SHA-256 does not match the source-locked original.')
    records = reconstruct(workbook)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=list(records[0]))
        writer.writeheader()
        writer.writerows(records)
    result = {
        'utility_provenance': 'New independent release utility; not historical SourceLock V3 code',
        'workbook_sha256': WORKBOOK_SHA256, 'generated_csv_sha256': sha256(output),
        'rows': len(records), 'columns': len(records[0]),
        'workbook_unchanged': sha256(workbook) == WORKBOOK_SHA256,
        'reference_used_to_construct_values': False,
    }
    success = result['workbook_unchanged']
    if reference:
        headers_a, a = load_csv(output)
        headers_b, b = load_csv(reference)
        differences = []
        numeric = categorical = 0
        maximum = 0.0
        structural = headers_a == headers_b and len(a) == len(b)
        if structural:
            for i, (x, y) in enumerate(zip(a, b), 1):
                for key in headers_a:
                    try:
                        if key in CATEGORICAL_FIELDS:
                            raise ValueError('raw text must match exactly')
                        u, v = float(x[key]), float(y[key])
                    except ValueError:
                        categorical += 1
                        same = x[key] == y[key]
                    else:
                        numeric += 1
                        finite = math.isfinite(u) and math.isfinite(v)
                        if finite:
                            maximum = max(maximum, abs(u - v))
                        same = finite and math.isclose(u, v, rel_tol=1e-12, abs_tol=1e-12)
                    if not same:
                        differences.append({'row': i, 'column': key,
                                            'generated': x[key], 'reference': y[key]})
        success = success and structural and not differences
        result.update({
            'reference_sha256': sha256(reference), 'same_structure': structural,
            'byte_identical_to_reference': sha256(output) == sha256(reference),
            'numeric_cells_compared': numeric, 'categorical_cells_compared': categorical,
            'maximum_absolute_numeric_difference': maximum,
            'numeric_absolute_tolerance': 1e-12, 'numeric_relative_tolerance': 1e-12,
            'mismatches': differences,
        })
    result['status'] = 'PASS' if success else 'FAIL'
    report.parent.mkdir(parents=True, exist_ok=True)
    report.write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f"{result['status']}: reconstructed {len(records)} rows; "
          f"byte-identical reference: {result.get('byte_identical_to_reference', 'not compared')}")
    return 0 if success else 1


if __name__ == '__main__':
    raise SystemExit(main())
